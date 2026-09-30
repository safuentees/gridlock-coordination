#!/usr/bin/env python3
"""Normalize the accepted public MISO Appendix A component inventory, without network access."""
import argparse
from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import re

import openpyxl

SOURCE = 'miso-mtep'
EDITION = 'appendix-a-2026-08-04'
LABEL = 'MISO Appendix A status through August 4, 2026'
GROUP = 'MISO MTEP inventory'
AS_OF = '2026-08-04'
RETRIEVED = '2026-09-30'
FILE = 'Appendix A Status Report575959.xlsx'
SHEET = 'MTEP Approved (Appendix A) ...'
URL = 'https://cdn.misoenergy.org/Appendix%20A%20Status%20Report575959.xlsx?v=20260514162226'
SHA = '55b8a60fc0a4159d9de52a5c1f83be04cfaa2c60468126408d9968782eb0c9dc'
LIFECYCLES = {'M1 - Proposed': 'planned', 'M2 - Appendix A Approved': 'approved',
              'M3 - Under Construction': 'under-construction'}
EXCLUDED = 'M4 - Project in Service'
EXCLUSION_REASON = 'M4 — completed and in service'
POLICY = json.loads(Path(__file__).with_name('policy.json').read_text())
SAMPLE = set(POLICY['sampleFacilityIds'])
CONTEXT = {item['nativeFacilityId']: item for item in POLICY['planningContextDecisions']}
OWNER_HOLDS = set(POLICY['ownerUnresolvedNormalizedTokens'])


def text(value):
    return None if value is None else str(value)


def native_id(value):
    if isinstance(value, (int, float)) and int(value) == value:
        return str(int(value))
    return str(value).strip() if value is not None else ''


def ref(locator):
    return {'file': FILE, 'locator': f"'{SHEET}'!{locator}", 'url': URL}


def expected_date(cell):
    result = {'value': None, 'precision': 'unknown', 'kind': 'in-service', 'label': 'Expected facility in-service date'}
    raw = cell.value
    if isinstance(raw, (date, datetime)):
        fmt = (cell.number_format or '').split(';')[0].lower()
        fmt = re.sub(r'"[^"]*"|\\.|\[[^]]*\]|_.|\*.', '', fmt)
        if 'y' in fmt:
            if 'm' in fmt and 'd' in fmt:
                result.update(value=raw.date().isoformat() if isinstance(raw, datetime) else raw.isoformat(), precision='day')
            elif 'm' in fmt:
                result.update(value=raw.strftime('%Y-%m'), precision='month')
            else:
                result.update(value=raw.strftime('%Y'), precision='year')
    # This pinned edition has only Excel dates. Unrecognized values are not guessed.
    return result


def owner_disposition(raw):
    key = re.sub(r'\s+', '', raw or '').upper()
    if not raw:
        return 'blank'
    if ',' in raw or '/' in raw or raw in {'MISO', 'To Be Determined', 'Local TO(s)'}:
        return 'joint-or-generic'
    if key in OWNER_HOLDS:
        return 'alias-needs-reconciliation'
    return 'source-token'


def owner(raw, row):
    if owner_disposition(raw) != 'source-token':
        return None
    slug = re.sub(r'[^a-z0-9]+', '-', re.sub(r'\s+', '', raw).lower()).strip('-')
    return {'id': 'miso-owner-' + slug, 'name': raw, 'identityScope': 'source', 'evidence': [ref(f'K{row}')]}


def kind(uid, facility_type):
    if uid in CONTEXT:
        return 'planning-context'
    if facility_type == 'Misc.':
        return 'capital-project-context'
    if facility_type in {'Substation', 'Transformer', 'Line New', 'Line Upgrade', 'Voltage Device'}:
        return 'transmission-project'
    return 'planning-context'


def project(cells, row):
    values = [cell.value for cell in cells]
    uid, pid = native_id(values[7]), native_id(values[4])
    status, facility_type = text(values[11]), text(values[14])
    record_kind = kind(uid, facility_type)
    listed_owner = text(values[10])
    qualified_owner = owner(listed_owner, row)
    timing = expected_date(cells[13])
    notes = [f'Source parent project {pid}: {values[5]}',
             'Submitting TO: ' + (text(values[2]) or 'not supplied'),
             f'Target cycle: {values[0]}; current appendix: {values[1]}; planning region: {values[3]}',
             f'Source State 1: {values[22]}; State 2: {values[23] if values[23] is not None else "not supplied"}']
    if values[17] is not None:
        notes.append('Source circuit ID: ' + text(values[17]))
    if values[12] is not None:
        notes.append('Source current estimated facility cost: ' + text(values[12]) + '; this is not a savings estimate.')
    miles = [label + ': ' + text(values[index]) for index, label in ((20, 'new'), (21, 'upgrade')) if values[index] is not None]
    if miles:
        notes.append('Source estimated miles, not route geometry or shared length: ' + '; '.join(miles))
    if values[24] is not None:
        notes.append('Source cost-shared status: ' + text(values[24]))
    if status == 'M1 - Proposed':
        notes.append('The facility is Proposed although parent Current Appendix is A. Its row status is preserved; approval is not inferred from the container or target cycle.')
    elif status == 'M2 - Appendix A Approved':
        notes.append('Appendix A approval includes engineering, right-of-way acquisition and procurement; it does not establish that construction has started.')
    elif status == 'M3 - Under Construction':
        notes.append('The source reports construction has begun; no construction start/end dates or interval are supplied.')
    else:
        notes.append('Unrecognized source status; lifecycle phase remains unknown.')
    if qualified_owner is None:
        reasons = {'blank': 'Facility Owner(s) is blank', 'joint-or-generic': 'the listed owner is joint or generic',
                   'alias-needs-reconciliation': 'the listed owner belongs to an unresolved alias family'}
        notes.append('Qualified owner identity is unknown because ' + reasons[owner_disposition(listed_owner)] + '. Submitting TO is not substituted.')
    sources = [ref(f'A{row}:Y{row}')]
    if uid in CONTEXT:
        decision = CONTEXT[uid]
        notes.append(decision['reason'])
        sources.append(decision['citation'])
    elif record_kind == 'capital-project-context':
        notes.append('Miscellaneous facility work is retained as capital context, not counted as a transmission-qualified project or a new corridor.')
    elif record_kind == 'planning-context':
        notes.append('The source facility type does not establish physical transmission scope; retain planning context.')
    if timing['precision'] == 'unknown' and cells[13].value is not None:
        notes.append(f'Unresolved Expected ISD: {text(cells[13].value)} (display format {cells[13].number_format!r})')
    voltage = sorted({float(value) for value in (values[18], values[19]) if isinstance(value, (int, float)) and value > 0})
    if any(isinstance(value, (int, float)) and value <= 0 for value in (values[18], values[19])):
        notes.append(f'Source Max kV: {values[18]}; Min kV: {values[19]}. Nonpositive placeholders are not normalized as voltages.')
    source_state = ' / '.join(text(v) for v in (values[22], values[23]) if v is not None) or None
    source_location = '; '.join(label + ': ' + text(v) for label, v in (('From Sub', values[15]), ('To Sub', values[16])) if v is not None) or None
    return {'id': f'{SOURCE}:{uid}', 'utilityId': SOURCE, 'sourceId': uid, 'sponsor': listed_owner,
            'name': text(values[8]) or uid, 'description': text(values[9]) or '', 'status': status,
            'voltageKv': voltage, 'date': timing, 'buildWindow': None, 'location': None,
            'sources': sources, 'notes': notes,
            'inventory': {'sourceKey': SOURCE, 'editionId': EDITION, 'editionLabel': LABEL,
                'reportingGroupName': GROUP, 'recordKind': record_kind, 'recordUnit': 'component',
                'identityBasis': 'native-id', 'lifecycle': LIFECYCLES.get(status, 'unknown'),
                'lifecycleAsOf': AS_OF, 'lifecycleDateMeaning': 'source-edition', 'owner': qualified_owner,
                'parentProgramId': f'{SOURCE}:project:{pid}', 'aliases': [],
                'workType': ' — '.join(text(v) for v in (values[6], facility_type) if v is not None) or None,
                'sourceState': source_state, 'sourceLocation': source_location, 'sourceUpdatedOn': None,
                'milestones': [{'field': 'Expected ISD', 'date': timing, 'evidence': [ref(f'N{row}')]}] if cells[13].value is not None else [],
                'constructionStart': None, 'nativeGeometry': None, 'centerRole': 'unknown',
                'locationAssumptions': ['Named substations and states are context only. The public workbook supplies no coordinates, datum or accuracy; no center or route is inferred.']}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--scope', choices=('sample', 'full'), default='sample')
    parser.add_argument('--generated-at', default=datetime.now(timezone.utc).isoformat())
    args = parser.parse_args()
    if hashlib.sha256(args.input.read_bytes()).hexdigest() != SHA:
        raise ValueError('Workbook differs from the accepted public edition; qualify changes before importing.')
    workbook = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    sheet = workbook[SHEET]  # Deliberately never enumerate workbook metadata sheets.
    projects, ledger, seen, parent_ids = [], [], set(), set()
    kinds, owners = Counter(), Counter()
    excluded = 0
    for row, cells in enumerate(sheet.iter_rows(min_row=4, max_col=25), 4):
        values = [cell.value for cell in cells]
        if not any(value is not None for value in values):
            continue
        uid, pid, status = native_id(values[7]), native_id(values[4]), text(values[11])
        if not uid or not pid or uid in seen:
            raise ValueError(f'Missing/duplicate Facility ID or missing parent at row {row}')
        seen.add(uid)
        entry = {'sheet': SHEET, 'sourceRow': row, 'nativeFacilityId': uid, 'parentProjectId': pid,
                 'rawStatus': status, 'rawFacilityType': text(values[14])}
        if args.scope == 'sample' and uid not in SAMPLE:
            entry.update(disposition='pending', reason='Outside the six-record connection sample.')
        elif status == EXCLUDED:
            excluded += 1
            entry.update(disposition='excluded', reason=EXCLUSION_REASON)
        else:
            record = project(cells, row)
            projects.append(record)
            parent_ids.add(pid)
            kinds[record['inventory']['recordKind']] += 1
            owners[owner_disposition(text(values[10]))] += 1
            entry.update(disposition='imported', reason='Nonterminal facility component retained.',
                         projectId=record['id'], recordKind=record['inventory']['recordKind'],
                         ownerDisposition=owner_disposition(text(values[10])))
        ledger.append(entry)
    workbook.close()
    pending = sum(row['disposition'] == 'pending' for row in ledger)
    counts = {'sourceRows': len(ledger), 'importedRecords': len(projects), 'excludedRows': excluded,
              'failedRows': 0, 'pendingRows': pending, 'locatedRecords': 0}
    manifest = {'sourceKey': SOURCE, 'editionId': EDITION, 'label': LABEL, 'reportingGroupId': SOURCE,
        'reportingGroupName': GROUP, 'publisher': 'Midcontinent Independent System Operator',
        'processAuthority': 'Midcontinent Independent System Operator', 'publishedOn': None, 'retrievedOn': RETRIEVED,
        'source': ref('A1:D1 report as-of; A4:Y3214 facility inventory'), 'sha256': SHA,
        'geographicScope': 'Partial MISO regional public facility inventory including seam work; state labels do not establish complete MISO or jurisdiction coverage.',
        'inventoryScope': 'All public Appendix A facility components except M4 in service. Explicit parent projects are not extra records. Ancillary and distribution context remain visible with distinct kinds; annual and under-evaluation inventories remain separate.',
        'scope': args.scope, 'counts': counts,
        'exclusions': [{'reason': EXCLUSION_REASON, 'count': excluded}] if excluded else [],
        'limitations': [
            'August 4, 2026 is the workbook report-through date, not a known publication or per-row update date. Retrieval is September 30; the URL cache parameter is not freshness.',
            'Facility status controls lifecycle. Target MTEP cycle and Current Appendix do not override a Proposed row; past expected dates do not prove completion.',
            'Miscellaneous capital work and the accepted distribution-context decisions are not transmission-qualified records. Parent IDs are context, not deduplicated engineering-project counts.',
            'Facility Owner(s) is separate from Submitting TO. Blank, joint, generic and unresolved alias families keep unknown qualified owners; other identities are source-scoped only.',
            'All coordinates, geometry and construction windows are unknown. Expected ISD is a facility forecast at displayed precision; source miles are not routes or shared corridors.',
            'Only the public visible inventory sheet is used; internal workbook metadata, portal models, maps and separate annual/evaluation inventories are not imported.',
            'Public access is not an open-data or commercial reuse license; source attribution and stated rights are preserved.',
        ]}
    dataset = {'schemaVersion': 2, 'mode': 'real', 'title': f'{LABEL} — {args.scope} inventory', 'generatedAt': args.generated_at,
        'projects': projects, 'sourceEditions': [manifest], 'coverage': {'processed': len(projects), 'skipped': excluded, 'failed': 0,
            'notes': [f'{len(ledger)} facility rows; {len(parent_ids)} retained parent project IDs; {pending} pending.',
                      'No native coordinates or source construction windows. Original workbook remains unchanged.']}}
    row_ledger = {'sourceKey': SOURCE, 'editionId': EDITION, 'sourceSha256': SHA, 'scope': args.scope,
        'counts': counts, 'retainedParentProjects': len(parent_ids), 'recordKindCounts': dict(kinds),
        'ownerDispositionCounts': dict(owners), 'rows': ledger}
    compact = json.dumps(dataset, ensure_ascii=False, separators=(',', ':'), allow_nan=False) + '\n'
    formatted = json.dumps(dataset, ensure_ascii=False, indent=2, allow_nan=False) + '\n'
    sizes = {'sourceKey': SOURCE, 'editionId': EDITION, 'scope': args.scope, 'records': len(projects),
        'compactUtf8Bytes': len(compact.encode('utf-8')), 'formattedDownloadUtf8Bytes': len(formatted.encode('utf-8')),
        'measurement': 'Source-only Dataset2 UTF-8 JSON; aggregate/export envelope can be larger.',
        'datasetFileFormat': 'compact' if args.scope == 'full' else 'formatted'}
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / 'dataset.json').write_text(compact if args.scope == 'full' else formatted)
    for filename, content in [('manifest.json', manifest), ('row-ledger.json', row_ledger), ('payload-sizes.json', sizes)]:
        (args.output_dir / filename).write_text(json.dumps(content, ensure_ascii=False, indent=2, allow_nan=False) + '\n')
    print(json.dumps({'counts': counts, 'parentProjects': len(parent_ids), 'recordKinds': dict(kinds),
                      'ownerDispositions': dict(owners), 'payloadSizes': sizes}))


if __name__ == '__main__':
    main()
