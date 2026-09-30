#!/usr/bin/env python3
"""Normalize the accepted public WestConnect TPPL without executing workbook macros."""
import argparse
from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

KEY = 'westconnect-tppl'
EDITION = '2026-02-10'
GROUP = 'WestConnect public TPPL'
LABEL = 'WestConnect public TPPL — February 10, 2026'
FILE = '02_10_26__wc_tppl_project_list_public.xlsm'
URL = 'https://doc.westconnect.com/Documents.aspx?NID=21174&dl=1'
SHA = 'ed9261738c81311d63afdd6c79b49d04f8e201477e01c0c9c08e5afe37570b3a'
FIELDS = ('projectid', 'Sponsor', 'ProjectName', 'Description', 'Purpose', 'Drivers', 'NeededForPublicPolicy_YorN', 'PublicPoliciesInvolved', 'FacilityType', 'Voltage', 'Development', 'Origin', 'Termination', 'InService', 'SubRegionalID_Name', 'Length', 'StateTraversed', 'WestConnect_TO_YorN', 'wcprojectID', 'Project_Map_URL', 'System_Impact_Study_YorN', 'All_Req_Infrastructure_Identifyied', 'Project_Inservice_for_Local_Planning_Process', 'Financing_Commitment', 'Financial_Arrangments', 'Construction_Commitments', 'Have_Project_Transmission_Service_Agreements', 'Transmission_Service_Agreement_Date', 'Federal_Permit_Required_YorN', 'Federal_Permit_ROD_YorN', 'State_Permit_Required_YorN', 'State_Permit_Obtained_YorN', 'Associated_Gen_YorN', 'Assoc_Gen_Cap', 'Assoc_Gen_Tech', 'Assoc_Gen_POI', 'Assoc_Gen_Sink', 'Assoc_Gen_TSA_YorN', 'Assoc_Gen_IA_YorN', 'Further_Details', 'modifieddate', 'ChangeTime', 'POST_DATE', 'DateProjectEntered')
COLUMNS = {name: get_column_letter(index + 1) for index, name in enumerate(FIELDS)}
EXPECTED_STATUS = {'Planned': 248, 'Conceptual': 56, 'Under Construction': 4, 'In-Service': 65, 'Withdrawn': 16}
TERMINAL = {'In-Service': 'explicit-in-service', 'Withdrawn': 'explicit-withdrawn'}
GENERATION_ONLY = {1575, 1579, 1580, 1581}
DISTRIBUTION_CONTEXT = {792, 1204, 1479, 1641}
PHYSICAL_OTHER = {1262, 1267, 1268, 1275, 1450, 1454, 1457, 1524, 1584, 1586, 1587, 1592, 1593, 1594, 1595, 1596, 1600, 1601, 1602, 1606, 1614, 1615, 1616, 1617, 1618, 1625, 1626, 1627, 1629, 1690, 1691}
SAMPLE_IDS = {5, 17, 1236, 1262, 1595, 1641}
CONTEXT_FIELDS = ('Purpose', 'Drivers', 'NeededForPublicPolicy_YorN', 'PublicPoliciesInvolved', 'SubRegionalID_Name', 'Length', 'WestConnect_TO_YorN') + FIELDS[20:40]


def clean(value):
    return ' '.join(str(value).split()) if value is not None else None


def raw_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def source(row=None, fields=()):
    locator = 'All Projects!A1:AR390; 389 project rows, header excluded'
    if row is not None:
        locator = f'All Projects!A{row}:AR{row}'
        if fields:
            locator += '; ' + ', '.join(f'{COLUMNS[field]}{row} {field}' for field in fields)
    return {'file': FILE, 'locator': locator, 'url': URL}


def exclusion(values):
    if values['Development'] in TERMINAL:
        return TERMINAL[values['Development']]
    return 'explicit-generation-only' if values['projectid'] in GENERATION_ONLY else None


def lifecycle(values):
    # Withdrawn is broader than cancelled and is never normalized to cancelled.
    return {'Planned': 'planned', 'Under Construction': 'under-construction', 'In-Service': 'in-service'}.get(values['Development'], 'unknown')


def record_kind(values):
    if values['projectid'] in DISTRIBUTION_CONTEXT:
        return 'planning-context'
    if values['FacilityType'] == 'Other' and values['projectid'] not in PHYSICAL_OTHER:
        raise ValueError(f"Unqualified Other facility: {values['projectid']}")
    if values['FacilityType'] not in ('Other', 'Substation', 'Transmission Line'):
        raise ValueError('Unqualified facility type; qualify changed source material separately.')
    return 'transmission-project'


def classification_basis(values):
    if values['projectid'] in DISTRIBUTION_CONTEXT:
        return 'explicit-distribution-substation-name-or-description'
    if values['projectid'] in PHYSICAL_OTHER:
        return 'accepted-explicit-physical-grid-equipment-set'
    return 'source-facility-type-and-project-scope'


def service_date(value):
    if value == 'TBD':
        return {'value': None, 'precision': 'unknown', 'kind': 'in-service', 'label': 'Expected in-service year'}
    if isinstance(value, int) and not isinstance(value, bool) and 1900 <= value <= 2100:
        return {'value': str(value), 'precision': 'year', 'kind': 'in-service', 'label': 'Expected in-service year'}
    raise ValueError(f'Unqualified expected in-service value: {value!r}')


def updated_day(value, epoch):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        value = from_excel(value, epoch)
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raise ValueError(f'Unqualified modifieddate value: {value!r}')


def voltage_values(values):
    texts = [values['ProjectName'] or '', values['Description'] or '']
    # A below-threshold category does not supply an exact voltage.
    category = values['Voltage'] or ''
    if not re.search(r'\bBelow\b', category, re.I):
        texts.append(category)
    numbers = set()
    for expression in re.findall(r'(?<![\w.])(\d+(?:\.\d+)?(?:\s*[/:-]\s*\d+(?:\.\d+)?)*)\s*kV\b', ' '.join(texts), re.I):
        numbers.update(float(number) for number in re.findall(r'\d+(?:\.\d+)?', expression))
    return sorted(number for number in numbers if number > 0)


def read_rows(path):
    if hashlib.sha256(path.read_bytes()).hexdigest() != SHA:
        raise ValueError('Input differs from the accepted public edition; qualify changed source material separately.')
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook['All Projects']
        if sheet.max_row != 390 or tuple(next(sheet.iter_rows(min_row=1, max_row=1, max_col=44, values_only=True))) != FIELDS:
            raise ValueError('Accepted project-table boundaries or headers changed.')
        rows = []
        # Read only the public project cells. No comments, contacts, form sheets,
        # hidden metadata, workbook properties or macros enter derived outputs.
        for number, cells in enumerate(sheet.iter_rows(min_row=2, max_row=390, max_col=44, values_only=True), 2):
            values = dict(zip(FIELDS, cells))
            native_id = values['projectid']
            if not isinstance(native_id, int) or isinstance(native_id, bool) or native_id <= 0:
                raise ValueError(f'Invalid native projectid at row {number}.')
            for field, value in values.items():
                if isinstance(value, str):
                    values[field] = clean(value)
            if not values['ProjectName'] or values['Development'] not in EXPECTED_STATUS:
                raise ValueError(f'Unqualified project name/lifecycle at row {number}.')
            rows.append({'row': number, 'values': values})
        if len({row['values']['projectid'] for row in rows}) != 389 or Counter(row['values']['Development'] for row in rows) != EXPECTED_STATUS:
            raise ValueError('The accepted 389-row identity/lifecycle enumeration changed.')
        return rows, workbook.epoch
    finally:
        workbook.close()


def make_project(row, epoch):
    number, values = row['row'], row['values']
    native_id, status = str(values['projectid']), values['Development']
    kind, milestone = record_kind(values), service_date(values['InService'])
    notes = [f"{field} ({COLUMNS[field]}{number}): {raw_value(values[field])}" for field in CONTEXT_FIELDS if values[field] is not None]
    notes.append(f"Source voltage category (J{number}): {values['Voltage']}.")
    combined_id = values['wcprojectID']
    if combined_id != f"{values['InService']}-{native_id}":
        notes.append(f"Source wcprojectID (S{number}) is {combined_id}; it differs from the current expected-year/projectid combination. Numeric projectid A{number} identifies this record and InService N{number} supplies its expected year.")
    if status == 'Conceptual':
        notes.append('Conceptual denotes active development that has not met the source planned criteria; it does not establish approval or construction.')
    if kind == 'planning-context':
        notes.append('The source explicitly describes a distribution substation; retain its planning context without asserting transmission-project scope.')
        if values['FacilityType'] == 'Transmission Line':
            notes.append('FacilityType says Transmission Line, but ProjectName and Description say distribution substation. Both source meanings are preserved.')
    if status == 'Under Construction' and isinstance(values['InService'], int) and values['InService'] < 2026:
        notes.append('The source still reports Under Construction with an elapsed expected year; the forecast does not establish completion.')
    for field in ('modifieddate', 'ChangeTime', 'POST_DATE', 'DateProjectEntered'):
        if values[field] is not None:
            notes.append(f"Source {field} ({COLUMNS[field]}{number}), raw value: {raw_value(values[field])}. Administrative timestamp; not a construction milestone or lifecycle-change date.")
    location_text = '; '.join(f'{field}: {values[field]}' for field in ('Origin', 'Termination') if values[field] is not None) or None
    return {
        'id': f'{KEY}:{native_id}', 'sourceId': native_id, 'utilityId': KEY, 'sponsor': values['Sponsor'],
        'name': values['ProjectName'], 'description': values['Description'] or '', 'status': status,
        'voltageKv': voltage_values(values), 'date': milestone, 'buildWindow': None, 'location': None,
        'sources': [source(number)], 'notes': notes,
        'inventory': {
            'sourceKey': KEY, 'editionId': EDITION, 'editionLabel': LABEL, 'reportingGroupName': GROUP,
            'recordUnit': 'project', 'identityBasis': 'native-id', 'recordKind': kind,
            'lifecycle': lifecycle(values), 'lifecycleAsOf': EDITION, 'lifecycleDateMeaning': 'source-edition',
            'sourceUpdatedOn': updated_day(values['modifieddate'], epoch), 'sourceState': values['StateTraversed'], 'sourceLocation': location_text,
            'owner': None, 'parentProgramId': None,
            'aliases': [f'Source wcprojectID: {combined_id}'] if combined_id else [],
            'workType': values['FacilityType'], 'constructionStart': None, 'nativeGeometry': None, 'centerRole': 'unknown',
            'locationAssumptions': ['Origin, termination and state are source text only. No native coordinates, geocoding, route, state allocation or center is supplied.'],
            'milestones': [{'field': 'InService', 'date': milestone, 'evidence': [source(number, ('InService',))]}],
        },
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--scope', choices=('sample', 'full'), default='sample')
    parser.add_argument('--generated-at', default=datetime.now(timezone.utc).isoformat())
    args = parser.parse_args()
    rows, epoch = read_rows(args.input)
    retained = [row for row in rows if not exclusion(row['values'])]
    if {row['values']['projectid'] for row in retained if row['values']['FacilityType'] == 'Other'} != PHYSICAL_OTHER:
        raise ValueError('The accepted physical Other classification set changed.')
    ledger, projects, exclusions = [], [], Counter()
    for row in rows:
        values = row['values']
        native_id, reason = str(values['projectid']), exclusion(values)
        entry = {'sourceId': native_id, 'row': row['row'], 'source': source(row['row']),
                 'sourceProjectName': values['ProjectName'], 'rawStatus': values['Development'],
                 'rawExpectedInService': values['InService'], 'rawCombinedProjectId': values['wcprojectID'],
                 'rawFacilityType': values['FacilityType'], 'rawVoltage': values['Voltage'],
                 'rawStateTraversed': values['StateTraversed'], 'rawOrigin': values['Origin'], 'rawTermination': values['Termination']}
        if reason:
            exclusions[reason] += 1
            entry.update(disposition='excluded', reason=reason)
            if reason == 'explicit-generation-only':
                entry['scopeEvidence'] = values['Description']
        else:
            entry.update(recordKind=record_kind(values), classificationBasis=classification_basis(values), lifecycle=lifecycle(values))
            if args.scope == 'full' or values['projectid'] in SAMPLE_IDS:
                project = make_project(row, epoch)
                projects.append(project)
                entry.update(disposition='imported', projectId=project['id'])
            else:
                entry.update(disposition='pending', reason='Accepted nonterminal row; normalization awaits sample connection.')
        ledger.append(entry)
    imported_count = 304 if args.scope == 'full' else 6
    counts = {'sourceRows': len(rows), 'importedRecords': len(projects), 'excludedRows': sum(exclusions.values()), 'failedRows': 0,
              'pendingRows': sum(entry['disposition'] == 'pending' for entry in ledger), 'locatedRecords': 0}
    if counts != {'sourceRows': 389, 'importedRecords': imported_count, 'excludedRows': 85, 'failedRows': 0, 'pendingRows': 304 - imported_count, 'locatedRecords': 0}:
        raise ValueError('Accepted source-scope reconciliation changed.')
    limitations = [
        'Public February 10, 2026 list only, with library posting March 26, 2026 separate from source edition. Actual publication date is unknown.',
        'This is partial WestConnect planning evidence, not a census of states or distinct engineering projects across other inventories.',
        'Numeric projectid is identity. Combined year/project labels include stale years and incomplete suffixes; they are only source references.',
        'Full accepted scope has 300 transmission-project and four explicit distribution-substation planning-context records; mixed scope remains visible.',
        'Sponsors identify source-reported owners or developers and do not establish individual asset ownership; owner identities remain unknown.',
        'All retained records are unlocated. Origin/termination/state text and approximate length do not supply coordinates, CRS, accuracy, routes or shared corridors.',
        'Planned does not establish regulatory approval; Conceptual remains unknown, and elapsed in-service forecasts do not establish completion.',
        'Withdrawn covers cancelled, postponed or no active development; withdrawn exclusions do not assert cancellation.',
        'Expected in-service dates have year precision or remain unknown. Administrative timestamps never become construction windows or comparison milestones.',
        'Factual source evidence is retained for authorized private analysis. No open-data license, original-workbook publication or unrestricted commercial reuse is asserted.',
    ]
    if args.scope == 'sample':
        limitations.append('Six retained diagnostic records are emitted first; 298 accepted records await provider connection and full normalization.')
    manifest = {
        'sourceKey': KEY, 'editionId': EDITION, 'label': LABEL, 'reportingGroupId': KEY, 'reportingGroupName': GROUP,
        'publisher': 'WestConnect', 'processAuthority': 'WestConnect regional transmission planning',
        'publishedOn': None, 'retrievedOn': '2026-09-30', 'source': source(), 'sha256': SHA,
        'geographicScope': 'Source-reported Arizona, New Mexico, Colorado, Texas, Wyoming, Nebraska, South Dakota, Multiple and blank state labels; partial regional submission scope, with no allocation of Multiple or inferred state coverage.',
        'inventoryScope': 'Public project-list entries with nonterminal transmission work and explicit distribution planning context. In-Service, Withdrawn and four exact generation-only entries are excluded.',
        'scope': args.scope, 'counts': counts, 'exclusions': [{'reason': reason, 'count': count} for reason, count in sorted(exclusions.items())],
        'limitations': limitations,
    }
    dataset = {'schemaVersion': 2, 'mode': 'real', 'title': LABEL + (' — six-record sample' if args.scope == 'sample' else ' — full accepted inventory'), 'generatedAt': args.generated_at,
               'projects': projects, 'sourceEditions': [manifest],
               'coverage': {'processed': len(projects), 'skipped': sum(exclusions.values()), 'failed': 0,
                            'notes': [f"All 389 source rows are reconciled; {counts['pendingRows']} accepted nonterminal records remain pending.", 'All records remain unlocated with unknown asset-owner identities.']}}
    mismatches = [row for row in rows if row['values']['wcprojectID'] != f"{row['values']['InService']}-{row['values']['projectid']}"]
    quality = {'sourceKey': KEY, 'editionId': EDITION, 'scope': args.scope,
               'fullSourceStatusCounts': dict(sorted(Counter(row['values']['Development'] for row in rows).items())),
               'acceptedRetainedLifecycleCounts': dict(sorted(Counter(lifecycle(row['values']) for row in retained).items())),
               'acceptedRetainedKindCounts': dict(sorted(Counter(record_kind(row['values']) for row in retained).items())),
               'acceptedRetainedDatePrecisionCounts': dict(sorted(Counter(service_date(row['values']['InService'])['precision'] for row in retained).items())),
               'importedLifecycleCounts': dict(sorted(Counter(project['inventory']['lifecycle'] for project in projects).items())),
               'importedKindCounts': dict(sorted(Counter(project['inventory']['recordKind'] for project in projects).items())),
               'combinedIdMismatchCounts': dict(sorted(Counter('missing-numeric-suffix' if row['values']['wcprojectID'].endswith('-') else 'different-year' for row in mismatches).items())),
               'combinedIdMismatches': [{'sourceId': str(row['values']['projectid']), 'source': source(row['row'], ('projectid', 'InService', 'wcprojectID')), 'rawCombinedId': row['values']['wcprojectID'], 'rawExpectedInService': row['values']['InService']} for row in mismatches],
               'classification': {'distributionContextIds': sorted(DISTRIBUTION_CONTEXT), 'physicalOtherIds': sorted(PHYSICAL_OTHER), 'generationOnlyExclusionIds': sorted(GENERATION_ONLY)},
               'sourceTimestampHandling': {'normalizedField': 'AO modifieddate', 'excelEpoch': epoch.isoformat(), 'otherFields': ['AP ChangeTime', 'AQ POST_DATE', 'AR DateProjectEntered'], 'meaning': 'Preserve raw value/locator; only modifieddate supplies sourceUpdatedOn. Other timestamps are labeled source context.'},
               'sampleIds': [f'{KEY}:{native_id}' for native_id in sorted(SAMPLE_IDS)]}
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for name, value in [('dataset.json', dataset), ('manifest.json', manifest), ('row-ledger.json', ledger), ('source-quality.json', quality)]:
        (args.output_dir / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')
    sizes = {'scope': args.scope, 'datasetRecords': len(projects),
             'compactUtf8Bytes': len(json.dumps(dataset, ensure_ascii=False, separators=(',', ':')).encode()),
             'formattedUtf8Bytes': (args.output_dir / 'dataset.json').stat().st_size}
    (args.output_dir / 'payload-sizes.json').write_text(json.dumps(sizes, indent=2) + '\n')
    print(json.dumps({'counts': counts, 'payloadSizes': sizes}))


if __name__ == '__main__':
    main()
