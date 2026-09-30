#!/usr/bin/env python3
"""Produce the accepted public ERCOT TPIT inventory without contact fields or inferred geometry."""
import argparse
from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re

import openpyxl
from openpyxl.utils import get_column_letter

KEY = 'ercot-tpit'
EDITION = '2026-07-13'
LABEL = 'ERCOT TPIT — July 13, 2026'
GROUP = 'ERCOT TPIT inventory'
FILE = 'ERCOT-July-Ad-Hoc-TPIT-No-Cost-071326-UPDATE.xlsx'
URL = 'https://www.ercot.com/files/docs/2022/03/02/' + FILE
SHA = '4a08ea0f26af780f0d63c21304f3969ee876a49b152b827be8b48b8edb655b05'
TABLES = ('FutureTPIT071326NoCost', 'PlannedTPIT071326NoCost', 'CompletedTPIT071326NoCost', 'CancelledTPIT071326NoCost')
SAMPLE_IDS = {'110733', '90446', '110174', '80422', '91626', '85979'}
EXPECTED_DUPLICATES = {'110733', '110749', '110751', '110753', '102795'}
# Exact single source tokens only: no corporate expansions, aliases or joint-owner splitting.
SINGLE_OWNER_TOKENS = {'ONCOR', 'CNP', 'AEP TCC', 'LCRATSC', 'AEP TNC', 'BEPC', 'STEC', 'TNMP', 'ETT', 'PEC', 'RCEC', 'CTT', 'BRYN', 'CPS', 'AEN', 'WETT', 'GVEC', 'LST', 'BBEC', 'BPUB', 'SLU', 'GPL', 'RGEC', 'GSEC', 'DME', 'BEC', 'NBU'}
PHYSICAL_TERMS = re.compile(r'\b(line|substation|switchyard|breaker|transformer|reactor|capacitor|rebuild|reconductor|upgrade)\w*\b', re.I)
DATE_COLUMNS = (('L', 'ProjectedInService', 'Projected in-service month/year', 'in-service'), ('M', 'ActualInService', 'Reported actual in-service month/year', 'in-service'), ('W', 'RPGSubmitted', 'Submitted to ERCOT for RPG review month/year', 'unknown'), ('X', 'RPGCompleted', 'RPG review completed month/year', 'unknown'), ('Y', 'BoardReviewCompleted', 'ERCOT board review completed month/year', 'unknown'))
RELATED_COLUMNS = {'H': 'Associated projects', 'K': 'Owner project number', 'V': 'RPG number', 'AE': 'Phase number', 'AF': 'MOD project number', 'AG': 'RTP project number'}


def source(locator):
    return {'file': FILE, 'locator': locator, 'url': URL}


def raw(value):
    return value.isoformat() if isinstance(value, (datetime, date)) else value


def text(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace('_x000D_', '\n').strip() or None


def finite_number(value):
    if isinstance(value, bool):
        return None
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (ValueError, TypeError):
        return None


def classification(values):
    supporting = [column for column in ('O', 'P', 'Q') if (finite_number(values[column]) or 0) > 0]
    reactive = values['R']
    if isinstance(reactive, (int, float)) and not isinstance(reactive, bool) and math.isfinite(reactive) and reactive != 0:
        supporting.append('R')
    if supporting:
        return 'physical-equipment', supporting
    scope_text = ' '.join(text(values[column]) or '' for column in ('B', 'C'))
    if re.search(r'\bload\s+records?\b', scope_text, re.I) and not PHYSICAL_TERMS.search(scope_text):
        return 'model-only', ['B', 'C']
    if re.search(r'\bdistribution\s+transformers?\b', scope_text, re.I):
        return 'distribution-context', ['B', 'C']
    return 'unresolved-context', ['B', 'C', 'O', 'P', 'Q', 'R']


def month_date(value, label, kind):
    if isinstance(value, (datetime, date)) and value.year != 9999:
        return {'value': f'{value.year:04d}-{value.month:02d}', 'precision': 'month', 'kind': kind, 'label': label}
    return {'value': None, 'precision': 'unknown', 'kind': kind, 'label': label}


def row_locator(row, columns='A:I; K:AG'):
    # The ranges intentionally omit contact column J.
    return f"{row['sheet']}!row {row['number']}, columns {columns}"


def read_rows(workbook):
    rows, by_id = [], {}
    duplicates = set()
    for sheet_name in TABLES:
        for number, cells in enumerate(workbook[sheet_name].iter_rows(min_row=3, max_col=33), 3):
            # Do not access J's value or any contact worksheet.
            values = {get_column_letter(index + 1): cell.value for index, cell in enumerate(cells) if index != 9}
            if values['A'] is None:
                continue
            native_id = text(values['A'])
            if not native_id:
                raise ValueError(f'Missing native ID at {sheet_name} row {number}')
            row = {'sheet': sheet_name, 'number': number, 'id': native_id, 'values': values, 'dateFormats': {column: cells[openpyxl.utils.column_index_from_string(column) - 1].number_format for column, *_ in DATE_COLUMNS}}
            if native_id in by_id:
                first = by_id[native_id]
                if values != first['values'] or native_id not in EXPECTED_DUPLICATES:
                    raise ValueError(f'Conflicting or unqualified repeated native ID {native_id}')
                row['duplicateOf'] = row_locator(first)
                first['duplicateLocators'].append(row_locator(row))
                duplicates.add(native_id)
            else:
                row['duplicateLocators'] = []
                by_id[native_id] = row
            rows.append(row)
    if len(rows) != 2127 or len(by_id) != 2122 or duplicates != EXPECTED_DUPLICATES:
        raise ValueError('The qualified TPIT enumeration changed; qualify the new source before importing.')
    return rows, by_id


def disposition(row):
    if 'duplicateOf' in row:
        return 'duplicate-row'
    if row['sheet'].startswith('Completed'):
        return 'completed-tab'
    if row['sheet'].startswith('Cancelled'):
        return 'cancelled-tab'
    status = text(row['values']['G'])
    if status == 'In-Service':
        return 'explicit-in-service'
    if status not in ('Planned', 'Under Construction', 'Conceptual', None):
        raise ValueError(f"Unqualified lifecycle at {row_locator(row)}")
    return 'model-only' if classification(row['values'])[0] == 'model-only' else None


def make_project(row):
    values, native_id = row['values'], row['id']
    category, supporting = classification(values)
    field_source = lambda column: source(f"{row['sheet']}!{column}{row['number']}; header {column}2")
    raw_status, owner_token = text(values['G']), text(values['I'])
    lifecycle = {'Planned': 'planned', 'Under Construction': 'under-construction'}.get(raw_status, 'unknown')
    owner = {'id': 'ercot-' + owner_token.lower().replace(' ', '-'), 'name': owner_token, 'identityScope': 'source', 'evidence': [field_source('I')]} if owner_token in SINGLE_OWNER_TOKENS else None
    physical = category == 'physical-equipment'
    classification_note = 'Physical transmission scope is supported by source equipment quantities: ' + ', '.join(f'{column}={text(values[column])}' for column in supporting) + '.' if physical else ('The source describes distribution-transformer context; physical transmission construction is not asserted.' if category == 'distribution-context' else 'Physical transmission scope is unresolved under the accepted equipment-evidence rule; missing or zero quantities do not prove that no physical work exists.')
    notes = [classification_note]
    if raw_status in ('Conceptual', None):
        notes.append(f"Source lifecycle is {raw_status or 'not supplied'}; normalized lifecycle remains unknown.")
    if owner is None:
        notes.append('The source owner token is unresolved or composite; it remains a label without a singular owner identity.')
    if text(values['D']):
        notes.append(f"Source comments: {text(values['D'])}")
    milestones = []
    for column, field, label, kind in DATE_COLUMNS:
        if values[column] is not None:
            milestone = month_date(values[column], label, kind)
            milestones.append({'field': field, 'date': milestone, 'evidence': [field_source(column)]})
            if milestone['precision'] == 'unknown':
                explanation = 'Year 9999 is an unresolved placeholder/implausible date, not usable event evidence.' if isinstance(values[column], (datetime, date)) and values[column].year == 9999 else 'The value does not establish a usable dated milestone.'
                notes.append(f"Unresolved {label}: raw {raw(values[column])} at {row['sheet']}!{column}{row['number']}. {explanation}")
    main_date = month_date(values['L'], 'Projected in-service month/year', 'in-service')
    if main_date['value'] and main_date['value'] < EDITION[:7]:
        notes.append('The projected in-service month predates this edition; it does not override the reported lifecycle or establish completion.')
    if values['M'] is not None:
        actual = month_date(values['M'], 'Reported actual in-service month/year', 'in-service')
        if actual['value']:
            notes.append('An actual in-service value is supplied on a nonterminal source row; the conflicting evidence is retained without inferring completion.')
        if actual['value'] and actual['value'] > EDITION[:7]:
            notes.append('The reported actual in-service month is later than the source edition. Its source label is preserved; it is not a confirmed energization date.')
    location_parts = [(label, text(values[column])) for column, label in [('E', 'From terminal'), ('F', 'To terminal'), ('S', 'Starting/substation county'), ('T', 'Ending county')]]
    voltage = finite_number(values['N'])
    return {
        'id': f'{KEY}:{native_id}', 'utilityId': KEY, 'sourceId': native_id, 'sponsor': owner_token,
        'name': text(values['B']) or f'TPIT record {native_id}', 'description': text(values['C']) or '', 'status': raw_status,
        'voltageKv': [voltage] if voltage is not None and voltage > 0 else [], 'date': main_date,
        'buildWindow': None, 'location': None,
        'sources': [source(row_locator(row))] + [source(locator) for locator in row['duplicateLocators']] + [field_source(column) for column in supporting], 'notes': notes,
        'inventory': {
            'sourceKey': KEY, 'editionId': EDITION, 'editionLabel': LABEL, 'reportingGroupName': GROUP,
            'recordKind': 'transmission-project' if physical else 'planning-context', 'recordUnit': 'component', 'identityBasis': 'native-id',
            'lifecycle': lifecycle, 'lifecycleAsOf': EDITION, 'lifecycleDateMeaning': 'source-edition', 'sourceUpdatedOn': None,
            'sourceState': None, 'sourceLocation': '; '.join(f'{label}: {value}' for label, value in location_parts if value) or None,
            'owner': owner, 'parentProgramId': None,
            'aliases': [f'{label}: {text(values[column])}' for column, label in RELATED_COLUMNS.items() if text(values[column])],
            'workType': None, 'constructionStart': None, 'nativeGeometry': None, 'centerRole': 'unknown',
            'locationAssumptions': ['The source provides named terminals/counties and bus references, not native coordinates. No center or route is inferred.'],
            'milestones': milestones,
        },
    }


def rtp_context(workbook, by_id):
    context = []
    for number, cells in enumerate(workbook['RTPProjects'].iter_rows(min_row=2, max_col=8, values_only=True), 2):
        if cells[0] is None:
            continue
        native_id = text(cells[0])
        explicit_links = re.findall(r'\d+[A-Za-z]?', text(cells[6]) or '')
        reverse_links = [row['id'] for row in by_id.values() if native_id in re.findall(r'\d{4}-[A-Za-z]+\d+', text(row['values']['AG']) or '')]
        context.append({'sourceRow': number, 'nativeId': native_id, 'source': source(f'RTPProjects!A{number}:H{number}'), 'rawFields': {get_column_letter(index + 1): raw(value) for index, value in enumerate(cells)}, 'explicitTpitIds': explicit_links, 'matchedTpitIds': [value for value in explicit_links if value in by_id], 'unresolvedTpitIds': [value for value in explicit_links if value not in by_id], 'reverseTpitIds': reverse_links, 'disposition': 'supplemental-context', 'reason': 'RTP membership can overlap TPIT; blank lifecycle and missing links do not establish an additional active project.'})
    if len(context) != 89:
        raise ValueError('Qualified RTP supplemental enumeration changed.')
    return context


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--scope', choices=('sample', 'full'), default='sample')
    parser.add_argument('--generated-at', default=datetime.now(timezone.utc).isoformat())
    args = parser.parse_args()
    if hashlib.sha256(args.input.read_bytes()).hexdigest() != SHA:
        raise ValueError('Input differs from the qualified public July 2026 edition; qualify the refresh separately.')
    workbook = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    rows, by_id = read_rows(workbook)
    ledger, projects, exclusions = [], [], Counter()
    for row in rows:
        values = row['values']
        category, supporting = classification(values)
        reason = disposition(row)
        entry = {'sourceSheet': row['sheet'], 'sourceRow': row['number'], 'nativeId': row['id'], 'locator': row_locator(row), 'rawStatus': raw(values['G']), 'rawOwnerToken': raw(values['I']), 'classification': category, 'classificationFields': supporting, 'rawQuantities': {column: raw(values[column]) for column in ('O', 'P', 'Q', 'R')}, 'rawScopeFields': {column: raw(values[column]) for column in ('B', 'C')}, 'rawMilestones': {column: raw(values[column]) for column, *_ in DATE_COLUMNS}, 'dateFormats': row['dateFormats']}
        if reason:
            entry.update(disposition='excluded', reason=reason)
            if 'duplicateOf' in row:
                entry['duplicateOf'] = row['duplicateOf']
            if reason in ('completed-tab', 'cancelled-tab'):
                entry['lifecycleEvidenceNote'] = 'Explicit source tab takes precedence; raw status and dates remain separate, potentially conflicting source assertions.'
            exclusions[reason] += 1
        elif args.scope == 'sample' and row['id'] not in SAMPLE_IDS:
            entry.update(disposition='pending', reason='Retained under accepted rules; normalization awaits first-sample connection.')
        else:
            projects.append(make_project(row))
            entry.update(disposition='imported', reason=category, projectId=f"{KEY}:{row['id']}")
        ledger.append(entry)
    supplemental = rtp_context(workbook, by_id)
    unusable_cells = []
    for row, entry in zip(rows, ledger):
        for column, *_ in DATE_COLUMNS:
            value = row['values'][column]
            if isinstance(value, (datetime, date)) and value.year == 9999:
                unusable_cells.append({'nativeId': row['id'], 'field': column, 'rawValue': raw(value), 'source': source(f"{row['sheet']}!{column}{row['number']}"), 'rowDisposition': entry['disposition']})
    quality = {'sourceKey': KEY, 'editionId': EDITION, 'scope': args.scope, 'rule': 'Year 9999 TPIT milestone cells are unknown/unusable, with exact raw evidence preserved.', 'unusableYear9999Cells': {'allSourceRowsIncludingDuplicates': len(unusable_cells), 'retainedDistinctRows': sum(cell['rowDisposition'] != 'excluded' for cell in unusable_cells), 'importedRows': sum(cell['rowDisposition'] == 'imported' for cell in unusable_cells)}, 'cells': unusable_cells}
    workbook.close()
    for item in supplemental:
        ledger.append({'sourceSheet': 'RTPProjects', 'sourceRow': item['sourceRow'], 'nativeId': item['nativeId'], 'locator': item['source']['locator'], 'rawStatus': item['rawFields']['F'], 'disposition': 'excluded', 'reason': 'supplemental-rtp-context'})
    exclusions['supplemental-rtp-context'] = len(supplemental)
    counts = {'sourceRows': len(ledger), 'importedRecords': len(projects), 'excludedRows': sum(exclusions.values()), 'failedRows': 0, 'pendingRows': sum(row['disposition'] == 'pending' for row in ledger), 'locatedRecords': 0}
    expected = {'duplicate-row': 5, 'completed-tab': 262, 'cancelled-tab': 78, 'explicit-in-service': 14, 'model-only': 2, 'supplemental-rtp-context': 89}
    if dict(exclusions) != expected or counts['sourceRows'] != 2216 or len(projects) != (6 if args.scope == 'sample' else 1766):
        raise ValueError('Accepted source dispositions changed; no output was written.')
    manifest = {
        'sourceKey': KEY, 'editionId': EDITION, 'label': LABEL, 'reportingGroupId': KEY, 'reportingGroupName': GROUP,
        'publisher': 'ERCOT', 'processAuthority': 'ERCOT transmission planning', 'publishedOn': '2026-07-17', 'retrievedOn': '2026-09-30',
        'source': source('Four TPIT project sheets, rows 3 onward (excluding contact column J); RTPProjects rows 2:90 supplemental only'), 'sha256': SHA,
        'geographicScope': 'ERCOT planning footprint in Texas; not all Texas transmission, and no per-record state or center is inferred.',
        'inventoryScope': 'Source-assigned TPIT component records: explicit transmission-equipment evidence and useful unresolved/distribution planning context. Completed/cancelled, explicit in-service and model-only entries excluded; RTP rows are supplemental context, not additional active projects.',
        'scope': args.scope, 'counts': counts, 'exclusions': [{'reason': reason, 'count': count} for reason, count in sorted(exclusions.items())],
        'limitations': ['All source records remain unlocated; named terminals/counties and bus references are not native geometry.', 'TPIT component counts are not unique engineering-project totals. Five exact non-contact duplicate rows are coalesced with both locators.', 'Only explicit accepted equipment quantities establish transmission-project scope; zero/missing quantities leave other useful records as planning-context.', 'Conceptual and absent status remain unknown; source tab/status and actual-date conflicts are preserved rather than resolved through date arithmetic.', 'Usable TPIT milestone fields retain displayed month precision; stored days/minutes are not day-level evidence. Year 9999 cells stay unknown with raw evidence preserved.', 'Owner tokens are source-scoped; unresolved/composite tokens do not become singular owner identities or globally distinct companies.', 'RTP records may overlap TPIT and have blank status; absent links do not establish uniqueness or cancellation.', 'Public raw data reused for analysis under ERCOT Terms of Use section 5; source attribution does not imply ERCOT endorsement.'],
    }
    dataset = {'schemaVersion': 2, 'mode': 'real', 'title': f'{LABEL} — {args.scope} inventory', 'generatedAt': args.generated_at, 'projects': projects, 'sourceEditions': [manifest], 'coverage': {'processed': len(projects), 'skipped': counts['excludedRows'], 'failed': 0, 'notes': ['Source-specific inventory components include qualified physical work and distinct planning context; no geometry or benefit is inferred.']}}
    formatted = json.dumps(dataset, ensure_ascii=False, indent=2, allow_nan=False) + '\n'
    compact = json.dumps(dataset, ensure_ascii=False, separators=(',', ':'), allow_nan=False)
    sizes = {'scope': args.scope, 'datasetRecords': len(projects), 'compactUtf8Bytes': len(compact.encode('utf-8')), 'formattedUtf8Bytes': len(formatted.encode('utf-8'))}
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / 'dataset.json').write_text(formatted)
    for filename, data in [('manifest.json', manifest), ('row-ledger.json', {'sourceKey': KEY, 'editionId': EDITION, 'sourceSha256': SHA, 'scope': args.scope, 'counts': counts, 'rows': ledger}), ('rtp-context.json', {'sourceKey': KEY, 'editionId': EDITION, 'note': 'Supplemental context only; never add these rows to active project counts.', 'rows': supplemental}), ('payload-sizes.json', sizes), ('source-quality.json', quality)]:
        (args.output_dir / filename).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + '\n')
    print(json.dumps({'counts': counts, 'sizes': sizes, 'unusableYear9999Cells': quality['unusableYear9999Cells'], 'ids': [project['sourceId'] for project in projects] if args.scope == 'sample' else 'full inventory'}, ensure_ascii=False))


if __name__ == '__main__':
    main()
