#!/usr/bin/env python3
"""Read a preserved public SPP workbook into the accepted C2 source contract."""
import argparse
from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import re

import openpyxl

SOURCE = 'spp-tracking'
EDITION = '2026-q3'
LABEL = 'SPP Q3 2026 project tracking'
PUBLISHED = '2026-07-07'
RETRIEVED = '2026-09-30'
FILE = 'Q3 2026 Quarterly Project Tracking Appendix 1.xlsm'
SHEET = 'Staff QPT Appendix 1'
SHA = '167bb301dc9fd290f7901318c43312b28513aec711484d1da17162a207f09b5f'
URL = 'https://www.spp.org/Documents/77416/3Q%202026%20Quarterly%20Project%20Tracking%20Report%20Appendix%201%262.zip'
SAMPLE = {'50442', '50444', '51487', '71933', '157811', '171264'}
EXCLUSIONS = {'Complete': 'Complete — construction complete and in service',
              'Closed Out': 'Closed Out — complete with closeout requirements fulfilled',
              'Withdrawn': 'Withdrawn — no longer retained in the active scope'}
LIFECYCLES = {'Delay - Mitigation': 'planned', 'Delay - Mitigation Window': 'planned',
              'On Schedule < 4': 'planned', 'On Schedule > 4': 'planned', 'Identified': 'planned',
              'NTC - Commitment Window': 'planned', 'NTC-C Project Estimate Window': 'planned',
              'In Service': 'under-construction', 'Re-evaluation': 'unknown',
              'Pending NTC/ NTC-C': 'unknown'}
CAVEATS = {
    'Delay - Mitigation': 'Source delay compares the owner forecast with the need date; mitigation is provided or time permits a change. Construction phase is not established.',
    'Delay - Mitigation Window': 'Source delay compares the owner forecast with the need date; mitigation is pending or time permits a change. Construction phase is not established.',
    'On Schedule < 4': 'Source schedule is within four years; it does not establish a construction interval or that construction has started.',
    'On Schedule > 4': 'Source schedule is beyond four years; it does not establish a construction interval or that construction has started.',
    'Identified': 'Identified in Board-approved study results; no construction commitment is inferred.',
    'NTC - Commitment Window': 'A notification has been issued, but the source status is within its written commitment window and no commitment has been received.',
    'NTC-C Project Estimate Window': 'Conditional notification estimate window; no construction commitment is inferred.',
    'In Service': 'Operational with construction activities still incomplete under the SPP definition; this is not a wholly completed upgrade.',
    'Re-evaluation': 'The source says active pending re-evaluation. Lifecycle stage remains unknown, not presumed suspended.',
    'Pending NTC/ NTC-C': 'The source glossary does not define this pending-notification status. Its lifecycle phase remains unknown.',
}
ALIAS_EVIDENCE = {
    'ekc': [
        {'file': 'SPP-PCWG-2021-06-02.pdf', 'locator': 'PDF p.109, Project Owner EKC/WR, UID 82116', 'url': 'https://www.spp.org/Documents/64759/PCWG%20Minutes%20and%20Attachments%2020210602.pdf'},
        {'file': 'SPP-EKC-name-change-2020.pdf', 'locator': 'PDF p.2, section II name change from Westar Energy to Evergy Kansas Central', 'url': 'https://www.spp.org/documents/62482/20200626_3220r1%20evergy%20ks%20central%20%26%20sunflower%20%26%20itc%20great%20plains%20interconnection%20agreement_er20-2187.pdf'},
    ],
    'em': [{'file': 'SPP-PCWG-2021-06-02.pdf', 'locator': 'PDF p.110, Project Owner EM/KCPL, UID 112370', 'url': 'https://www.spp.org/Documents/64759/PCWG%20Minutes%20and%20Attachments%2020210602.pdf'}],
    'emw': [{'file': 'SPP-EMW-name-change-2020.pdf', 'locator': 'PDF p.11, recital A, EMW formerly KCP&L Greater Missouri Operations (GMO)', 'url': 'https://www.spp.org/documents/62499/20200630_3024r1%20evergy%20missouri%20west%20and%20transource%20missouri%20interconnection%20agreement_er20-2278.pdf'}],
}
ALIASES = {'EKC': 'ekc', 'WR': 'ekc', 'EM': 'em', 'KCPL': 'em', 'EMW': 'emw', 'GMO': 'emw'}
ALIAS_NAMES = {'ekc': 'EKC / WR', 'em': 'EM / KCPL', 'emw': 'EMW / GMO'}
DATE_FIELDS = [('I', 'Project Owner Indicated In-Service Date', 'in-service'),
               ('J', 'RTO Determined Need Date', 'need'),
               ('K', 'Letter of Notification to Construct Issue Date', 'unknown')]


def text(value):
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def normalized(value):
    return ' '.join(str(value).split()) if value is not None else ''


def source_ref(locator):
    return {'file': FILE, 'locator': f"'{SHEET}'!{locator}", 'url': URL}


def date_value(cell, kind, label):
    """Respect visible calendar precision, not hidden Excel serial day values."""
    raw = cell.value
    value, precision = None, 'unknown'
    if isinstance(raw, (date, datetime)):
        # Remove literal text, escaped characters, color/locale and padding tokens.
        fmt = (cell.number_format or '').split(';')[0].lower()
        fmt = re.sub(r'"[^"]*"|\\.|\[[^]]*\]|_.|\*.', '', fmt)
        if 'y' in fmt:
            if 'm' in fmt and 'd' in fmt:
                value, precision = raw.date().isoformat() if isinstance(raw, datetime) else raw.isoformat(), 'day'
            elif 'm' in fmt:
                value, precision = raw.strftime('%Y-%m'), 'month'
            else:
                value, precision = raw.strftime('%Y'), 'year'
    elif isinstance(raw, str):
        displayed = raw.strip()
        if re.fullmatch(r'\d{4}', displayed):
            value, precision = displayed, 'year'
        elif re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', displayed):
            value, precision = displayed, 'month'
        else:
            for fmt in ('%m/%d/%Y', '%m-%d-%Y', '%Y-%m-%d'):
                try:
                    value, precision = datetime.strptime(displayed, fmt).date().isoformat(), 'day'
                    break
                except ValueError:
                    pass
    return {'value': value, 'precision': precision, 'kind': kind, 'label': label}


def owner(raw, row):
    token = normalized(raw)
    if not token or token.upper() in {'TBD', 'N/A', 'NA', 'UNKNOWN'} or re.search(r'[/|;&]', token):
        return None
    slug = ALIASES.get(token, re.sub(r'[^a-z0-9]+', '-', token.lower()).strip('-'))
    if not slug:
        return None
    return {'id': 'spp-owner-' + slug, 'name': ALIAS_NAMES.get(slug, token), 'identityScope': 'source',
            'evidence': [source_ref(f'D{row}')] + ALIAS_EVIDENCE.get(slug, [])}


def project(cells, row):
    vals = [cell.value for cell in cells]
    uid, pid, status = normalized(vals[2]), normalized(vals[1]), normalized(vals[17])
    listed_owner = text(vals[3])
    notes = [f'Source project PID {pid}: {vals[5]}', f'Source study: {vals[11]}',
             CAVEATS.get(status, 'Unrecognized source status; lifecycle phase remains unknown.')]
    milestones = []
    dates = []
    for index, (column, field, kind) in enumerate(DATE_FIELDS, 8):
        cell = cells[index]
        parsed = date_value(cell, kind, field)
        dates.append(parsed)
        if cell.value is not None:
            milestones.append({'field': field, 'date': parsed, 'evidence': [source_ref(f'{column}{row}')]})
            if parsed['precision'] == 'unknown':
                notes.append(f'Unresolved source {field}: {text(cell.value)} (display format {cell.number_format!r})')
    comparison = dates[0] if dates[0]['value'] else dates[1] if dates[1]['value'] else {'value': None, 'precision': 'unknown', 'kind': 'unknown', 'label': 'No usable owner forecast or RTO need date'}
    qualified_owner = owner(listed_owner, row)
    if qualified_owner is None:
        notes.append('Source ProjectOwner does not establish one owner: ' + (listed_owner or 'not supplied'))
    buses = []
    for index, label in ((18, 'From'), (20, 'To')):
        if vals[index] is not None or vals[index+1] is not None:
            buses.append(f'{label} bus {text(vals[index]) or "unknown number"}: {text(vals[index+1]) or "unknown name"}')
    cost_labels = [(12, 'baseline estimate'), (13, 'baseline estimate year'), (14, 'baseline estimate with escalation'), (15, 'current estimate'), (16, 'final cost')]
    costs = [f'{label}: {text(vals[index])}' for index, label in cost_labels if vals[index] is not None]
    if costs:
        notes.append('Raw source cost fields, not savings: ' + '; '.join(costs) + '. Zero/formula placeholders do not establish a free project.')
    counts = [(24, 'Number of New'), (25, 'Number of Rebuild/Reconductor'), (26, 'Number of Voltage Conversion')]
    quantities = [f'{label}: {text(vals[index])}' for index, label in counts if vals[index] is not None]
    if quantities:
        notes.append('Raw source quantity columns (units not specified in the headers; not geometry): ' + '; '.join(quantities))
    voltage_text = text(vals[23]) or ''
    voltage = sorted({float(v) for v in re.findall(r'\d+(?:\.\d+)?', voltage_text) if float(v) > 0})
    if voltage_text and (not voltage or re.search(r'(^|/)0($|/)', voltage_text)):
        notes.append('Source voltage field: ' + voltage_text + '; only positive voltage values are normalized.')
    return {'id': f'{SOURCE}:{uid}', 'utilityId': SOURCE, 'sourceId': uid, 'sponsor': listed_owner,
            'name': text(vals[6]) or uid, 'description': text(vals[22]) or '', 'status': text(vals[17]),
            'voltageKv': voltage, 'date': comparison, 'buildWindow': None, 'location': None,
            'sources': [source_ref(f'A{row}:AA{row}')], 'notes': notes,
            'inventory': {'sourceKey': SOURCE, 'editionId': EDITION, 'editionLabel': LABEL,
                'reportingGroupName': 'SPP project tracking', 'recordKind': 'transmission-project',
                'recordUnit': 'upgrade', 'lifecycle': LIFECYCLES.get(status, 'unknown'),
                'lifecycleAsOf': PUBLISHED, 'lifecycleDateMeaning': 'source-edition',
                'owner': qualified_owner, 'parentProgramId': f'{SOURCE}:pid:{pid}',
                'aliases': ['ntc:' + normalized(vals[0])] if vals[0] is not None else [],
                'workType': text(vals[7]), 'sourceState': text(vals[4]),
                'sourceLocation': '; '.join(buses) or None, 'sourceUpdatedOn': None,
                'milestones': milestones, 'constructionStart': None, 'nativeGeometry': None,
                'centerRole': 'unknown', 'locationAssumptions': ['The workbook provides named buses without coordinates or CRS; no center or route is inferred.']}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--scope', choices=('sample', 'full'), default='sample')
    parser.add_argument('--generated-at', default=datetime.now(timezone.utc).isoformat())
    args = parser.parse_args()
    if hashlib.sha256(args.input.read_bytes()).hexdigest() != SHA:
        raise ValueError('Workbook differs from the accepted Q3 2026 source; qualify changes before importing.')
    workbook = openpyxl.load_workbook(args.input, read_only=True, data_only=True)
    sheet = workbook[SHEET]
    projects, ledger, seen, parents = [], [], set(), set()
    excluded = Counter()
    for row, cells in enumerate(sheet.iter_rows(min_row=3), 3):
        vals = [cell.value for cell in cells]
        if not any(value is not None for value in vals):
            continue
        uid, pid, status = normalized(vals[2]), normalized(vals[1]), normalized(vals[17])
        if not uid or not pid or uid in seen:
            raise ValueError(f'Missing/repeated UID or absent parent PID at row {row}')
        seen.add(uid)
        entry = {'sheet': SHEET, 'sourceRow': row, 'nativeId': uid, 'parentPid': pid,
                 'rawStatus': text(vals[17]), 'normalizedStatus': status}
        if args.scope == 'sample' and uid not in SAMPLE:
            entry.update(disposition='pending', reason='Outside the six-record connection sample.')
        elif status in EXCLUSIONS:
            reason = EXCLUSIONS[status]
            excluded[reason] += 1
            entry.update(disposition='excluded', reason=reason)
        else:
            record = project(cells, row)
            projects.append(record)
            parents.add(pid)
            entry.update(disposition='imported', reason='Nonterminal transmission upgrade retained.', projectId=record['id'])
        ledger.append(entry)
    workbook.close()
    pending = sum(row['disposition'] == 'pending' for row in ledger)
    counts = {'sourceRows': len(ledger), 'importedRecords': len(projects), 'excludedRows': sum(excluded.values()),
              'failedRows': 0, 'pendingRows': pending, 'locatedRecords': 0}
    manifest = {'sourceKey': SOURCE, 'editionId': EDITION, 'label': LABEL, 'reportingGroupId': SOURCE,
        'reportingGroupName': 'SPP project tracking', 'publisher': 'Southwest Power Pool',
        'processAuthority': 'Southwest Power Pool', 'publishedOn': PUBLISHED, 'retrievedOn': RETRIEVED,
        'source': source_ref('A3:AA1424; complete UID inventory'), 'sha256': SHA,
        'geographicScope': 'Partial SPP regional public transmission-upgrade portfolio; raw states and named buses do not establish complete jurisdiction coverage.',
        'inventoryScope': 'All Appendix 1 UID upgrades except Complete, Closed Out and Withdrawn. Parent PIDs are relationships, not additional inventory records; Appendix 2 adds no independent upgrades.',
        'scope': args.scope, 'counts': counts,
        'exclusions': [{'reason': reason, 'count': count} for reason, count in sorted(excluded.items())],
        'limitations': [
            'Report publication/lifecycle edition anchor is July 7, 2026; the archive index was posted August 5, 2026. A conflicting older reporting-period string does not establish an activity interval.',
            'Every location is unknown. Bus numbers/names and state labels are raw context, not coordinates or routes.',
            'Owner identities are source-scoped; EKC/WR, EM/KCPL and EMW/GMO are evidenced aliases. TBD/blank/composite owners remain unknown; no ultimate parent is inferred.',
            'In Service means operational with incomplete construction. Schedule/notification stages do not establish construction start; re-evaluation and undefined pending notifications retain unknown phase.',
            'Owner forecast, RTO need and notification dates remain separate at displayed source precision; none creates a construction window.',
            'Appendix 2 contains 26 UIDs already in Appendix 1 and is deferred; its inconsistent free-text histories do not overwrite forecasts.',
            'SPP terms distinguish cited copying from commercial use requiring express authorization. This attributed analysis does not claim commercial permission or an open-data license.',
        ]}
    dataset = {'schemaVersion': 2, 'mode': 'real', 'title': f'{LABEL} — {args.scope} inventory',
        'generatedAt': args.generated_at, 'projects': projects, 'sourceEditions': [manifest],
        'coverage': {'processed': len(projects), 'skipped': sum(excluded.values()), 'failed': 0,
            'notes': [f'{len(ledger)} UID rows; {len(parents)} distinct retained parent PIDs; {pending} pending.',
                      'No source coordinates. Original workbook remains unchanged and no macros are executed.']}}
    ledger_output = {'sourceKey': SOURCE, 'editionId': EDITION, 'sourceSha256': SHA, 'scope': args.scope,
                     'counts': counts, 'retainedParentPids': len(parents), 'rows': ledger}
    compact = json.dumps(dataset, ensure_ascii=False, separators=(',', ':'), allow_nan=False) + '\n'
    formatted = json.dumps(dataset, ensure_ascii=False, indent=2, allow_nan=False) + '\n'
    sizes = {'sourceKey': SOURCE, 'editionId': EDITION, 'scope': args.scope, 'records': len(projects),
             'compactUtf8Bytes': len(compact.encode('utf-8')), 'formattedDownloadUtf8Bytes': len(formatted.encode('utf-8')),
             'measurement': 'Source-only Dataset2 UTF-8 JSON; aggregate/export envelope can be larger.',
             'datasetFileFormat': 'compact' if args.scope == 'full' else 'formatted'}
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / 'dataset.json').write_text(compact if args.scope == 'full' else formatted)
    for filename, content in [('manifest.json', manifest), ('row-ledger.json', ledger_output), ('payload-sizes.json', sizes)]:
        (args.output_dir / filename).write_text(json.dumps(content, ensure_ascii=False, indent=2, allow_nan=False) + '\n')
    print(json.dumps({'counts': counts, 'retainedParentPids': len(parents), 'payloadSizes': sizes}))


if __name__ == '__main__':
    main()
