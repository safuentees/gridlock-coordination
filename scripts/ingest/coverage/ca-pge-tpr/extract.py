#!/usr/bin/env python3
"""Normalize the accepted public May 2026 PG&E TPR, without restricted fields."""
import argparse
from collections import Counter
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

KEY = 'ca-pge-tpr'
EDITION = '2026-05-01'
GROUP = 'PG&E public TPR'
LABEL = 'PG&E public May 2026 TPR portfolio'
FILE = 'pge-may2026-public.xlsx'
URL = 'https://pgera.azurewebsites.net/Regulation/ValidateDocAccess?docID=871875'
SHA = '917464588f082e2239726479b331dff6ffda97436af9daa5372332a734d4b94e'
LETTER_FILE = 'pge-may2026-transmittal.pdf'
LETTER_URL = 'https://pgera.azurewebsites.net/Regulation/ValidateDocAccess?docID=871864'
LETTER_SHA = '50b35533feee7eecb1c8da83269716390227559bb74412a558aeed3f62937393'
SHEET = 'CPUC TPR PUBLIC'
SAMPLE_IDS = {'5771718', '5792705', '5779902', '5562465', 'Ex112888', '5816828'}
REPLACEMENT_NAMES = (
    'Electric Transmission Substation Security Upgrade',
    'Electric Transmission Cyber Security Upgrade',
    'Work at the Request of Others (WRO) Network Upgrades',
    'Private Load Interconnection Project',
)
EXPECTED_STATUS = {
    'Operational': 1079, 'Cancelled': 1, 'Planning': 390,
    'Engineering less than 50% complete': 234, 'Engineering more than 50% complete': 521,
    'Programmatic Work': 185, 'On Hold': 123,
    'Construction (under 10%)': 9, 'Construction (10%-25%)': 24,
    'Construction (25%-50%)': 17, 'Construction (50%-75%)': 37, 'Construction (over 75%)': 51,
}
EXPECTED_KIND = {'transmission-project': 1001, 'capital-project-context': 278, 'planning-context': 312}
EXPECTED_LIFECYCLE = {'planned': 1145, 'under-construction': 138, 'on-hold': 123, 'unknown': 185}
# Selected public, non-person fields only. Y (Project Manager) is never retained.
COLUMNS = tuple('B C D E F G H I J K L M N O P Q R S T U V W X Z AA AB AC AI AJ AK AL AM AN AO AP AQ AR AS AT AU AV AW AX AY AZ BA BB BC BD BE BF BG BH BI BK BL BM BN BO BP BQ BR BS BT BU BV BW BX BY BZ CA CB CC CD CE CF CG CH CI CJ CK CL CM CN CO CP CQ CS CT CU CV'.split())
EXPECTED_HEADERS = {
    'B': 'Row/Line No.', 'C': 'Project Name(s)', 'D': 'Latitude', 'E': 'Longitude',
    'F': 'Location2', 'G': 'Project Description', 'H': 'Project Description - What',
    'AI': 'Utility Unique ID #1 (Most Specific)', 'AJ': 'Utility Unique ID #2 (Less Specific)',
    'BI': 'Project Status', 'BK': 'Construction Start Date',
    'BL': 'Original Planned In-Service Date', 'BM': 'Current Projected or Actual In-Service Date',
}
CONTEXT_GROUPS = (
    ('Scope and purpose', 'H I J K L M N O P Q R S T U V W X Z AA'.split()),
    ('Source identity and approval', 'AJ AK AL AM AN AO AP AQ AR AS AT AU AV AW AX'.split()),
    ('Permit evidence', 'AY AZ BA BB BC BD BE BF BG BH'.split()),
    ('Service-date and rate-base context', 'BN BO BP'.split()),
    ('Source costs and expenditures ($000)', 'BQ BR BS BT BU BV BW BX BY BZ CA CB CC CD CE CF CG CI CJ CK CL CM CN'.split()),
    ('Other source financial context', 'CH CO CP CQ CS CT CU'.split()),
)


def text(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def meaningful(value):
    return text(value) not in (None, '', 'NA', 'N/A')


def source(row=None, columns=()):
    locator = f'{SHEET}!B5:CV2675; all 2671 native-ID rows; manager column Y omitted'
    if row is not None:
        locator = f'{SHEET}!B{row}:X{row}, Z{row}:CV{row}'
        if columns:
            locator += '; ' + ', '.join(f'{column}{row}' for column in columns)
    return {'file': FILE, 'locator': locator, 'url': URL}


def letter_source():
    return {'file': LETTER_FILE, 'url': LETTER_URL, 'locator': 'Pages 2–3, 5, 7–8; dates, native units, programmatic work and public redactions'}


def definition_source(rows):
    return {'file': FILE, 'url': URL, 'locator': f'Data Field Descriptions!B{rows[0]}:F{rows[-1]}'}


def exclusion(values):
    return {'Operational': 'explicit-operational', 'Cancelled': 'explicit-cancelled'}.get(values['BI'])


def kind_and_basis(values):
    description = (text(values['G']) or '').lower()
    name = (text(values['C']) or '').lower()
    native_id = text(values['AI'])
    equipment = text(values['H']) or ''
    if ('forecast placeholder' in description or native_id.lower().startswith('ex')
            or 'redact' in description or any(pattern.lower() in name for pattern in REPLACEMENT_NAMES)):
        return 'planning-context', 'accepted-forecast-investment-or-public-redaction-rule'
    if ('bucket for' in description or values['BI'] == 'Programmatic Work'
            or equipment == 'Transmission Line - Right of Way'
            or not equipment.startswith(('Substation -', 'Transmission Line -', 'Transmission Line/Substation'))):
        return 'capital-project-context', 'accepted-capital-program-budget-or-supporting-scope-rule'
    return 'transmission-project', 'explicit-physical-equipment-class-in-ferc-transmission-portfolio'


def lifecycle(status):
    if status == 'Programmatic Work':
        return 'unknown'
    if status == 'On Hold':
        return 'on-hold'
    if status.startswith('Construction ('):
        return 'under-construction'
    if status in ('Planning', 'Engineering less than 50% complete', 'Engineering more than 50% complete'):
        return 'planned'
    raise ValueError(f'Unqualified retained lifecycle: {status!r}')


def milestone(value, label, kind='in-service'):
    if isinstance(value, (datetime, date)):
        return {'value': text(value), 'precision': 'day', 'kind': kind, 'label': label}
    if value not in (None, 'NA', 'TBD', 'OAI'):
        raise ValueError(f'Unqualified date cell: {value!r}')
    suffix = 'OAI (operative as installed; no event date)' if value == 'OAI' else text(value) or 'not supplied'
    return {'value': None, 'precision': 'unknown', 'kind': 'unknown', 'label': f'{label} — {suffix}'}


def voltages(values):
    found = set()
    line = text(values['AB'])
    if meaningful(line):
        if not re.fullmatch(r'\d+(?:\.\d+)?(?:\s*\|\s*\d+(?:\.\d+)?)*', line):
            raise ValueError(f'Unqualified line-voltage expression: {line!r}')
        found.update(float(item.strip()) for item in line.split('|'))
    station = text(values['AC'])
    if meaningful(station):
        # This frozen edition has explicit kV strings only; never convert MVA.
        if not re.fullmatch(r'[\d. |]+\s*kV', station, re.I):
            raise ValueError(f'Unqualified substation-capacity units: {station!r}')
        found.update(float(item.strip()) for item in re.sub(r'\s*kV$', '', station, flags=re.I).split('|'))
    if any(number <= 0 for number in found):
        raise ValueError('Voltage must be positive when explicitly supplied.')
    return sorted(found)


def parent_id(values):
    value = text(values['AJ'])
    return value if value not in (None, '', 'NA', 'TBD') else None


def read_rows(original_dir):
    for name, expected_hash in ((FILE, SHA), (LETTER_FILE, LETTER_SHA)):
        if hashlib.sha256((original_dir / name).read_bytes()).hexdigest() != expected_hash:
            raise ValueError(f'Original differs from the accepted public edition: {name}')
    workbook = load_workbook(original_dir / FILE, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook[SHEET]
        if sheet.sheet_state != 'visible' or sheet.max_row != 2675:
            raise ValueError('Accepted visible project-table boundary changed.')
        header_values = next(sheet.iter_rows(min_row=4, max_row=4, max_col=100, values_only=True))
        headers = {column: text(header_values[column_index_from_string(column) - 1]) for column in COLUMNS}
        if any(headers[column] != expected for column, expected in EXPECTED_HEADERS.items()):
            raise ValueError('Accepted project-table field meanings changed.')
        rows = []
        for number, cells in enumerate(sheet.iter_rows(min_row=5, max_row=2675, max_col=100, values_only=True), 5):
            values = {column: cells[column_index_from_string(column) - 1] for column in COLUMNS}
            # Neither Y, comments, hidden metadata, other sheets nor restricted copies
            # are retained in the source-row objects or derived output.
            native_id = text(values['AI'])
            if not native_id or values['B'] != number - 4 or not text(values['C']):
                raise ValueError(f'Invalid native identity/locator at row {number}.')
            if values['D'] is not None or values['E'] is not None:
                raise ValueError('Accepted public coordinate fields are no longer all blank; requalify location meaning.')
            rows.append({'row': number, 'values': values})
        if len({text(row['values']['AI']) for row in rows}) != 2671 or Counter(row['values']['BI'] for row in rows) != EXPECTED_STATUS:
            raise ValueError('Accepted 2671-row identity/status enumeration changed.')
        return rows, headers
    finally:
        workbook.close()


def make_project(row, headers):
    number, values = row['row'], row['values']
    native_id, status = text(values['AI']), values['BI']
    kind, _ = kind_and_basis(values)
    parent = parent_id(values)
    current = milestone(values['BM'], 'Current projected or actual in-service')
    original = milestone(values['BL'], 'Original planned in-service')
    start = milestone(values['BK'], 'Actual or expected construction start', kind='unknown')
    notes = []
    for label, columns in CONTEXT_GROUPS:
        parts = [f'{headers[column]} ({column}{number}): {text(values[column])}' for column in columns if meaningful(values[column])]
        if parts:
            notes.append(f'{label}: ' + '; '.join(parts))
    notes.append(f'Source voltage evidence: AB{number} line kV = {text(values["AB"])}; AC{number} substation/transformer capacity = {text(values["AC"])}. Values retain their separate source field meanings.')
    if meaningful(values['CV']):
        notes.append(f'Source notes (CV{number}): {text(values["CV"])}')
    if status.startswith(('Engineering', 'Construction')):
        notes.append('The source derives engineering/construction percentage bands from milestone timing, not measured physical completion; retain the reported status with that limitation.')
    if status == 'Programmatic Work':
        notes.append('Programmatic work is an allocation or bucket, not an individually identified engineering project; lifecycle remains unknown.')
    if kind == 'planning-context':
        notes.append('Public forecast, investment-code or redacted scope does not establish a defined physical project. Restricted underlying details are not reconstructed.')
    if parent:
        notes.append('The supplied ID #2 groups source work; the parent description/cost may span multiple planning orders and owners. No parent record, unique engineering-project count or singular asset owner is inferred.')
    if meaningful(values['BQ']):
        notes.append('Original projected cost/range is the ID #2 project total where available. Source costs are $000 and are not additive across repeated parent entries or evidence of coordination savings.')
    else:
        notes.append('Source cost fields retain $000 units and their own forecast/actual meanings; they do not establish coordination savings.')
    notes.append('March 2026 status evidence is separate from the March 10 automated extraction and May 1 publication. Actual-or-expected start and projected-or-actual service dates do not establish a construction window.')
    inventory = {
        'sourceKey': KEY, 'editionId': EDITION, 'editionLabel': LABEL, 'reportingGroupName': GROUP,
        'recordKind': kind, 'identityBasis': 'native-id', 'lifecycle': lifecycle(status),
        'lifecycleAsOf': '2026-03', 'lifecycleDateMeaning': 'source-edition',
        'sourceUpdatedOn': None, 'sourceState': None, 'sourceLocation': text(values['F']),
        'owner': None, 'parentProgramId': parent,
        'aliases': [f'{label}: {text(values[column])}' for column, label in (('AK', 'Source ID #2 name'), ('AM', 'Major work category')) if meaningful(values[column])],
        'workType': text(values['H']), 'constructionStart': None,
        'nativeGeometry': None, 'centerRole': 'unknown',
        'locationAssumptions': ['All public latitude/longitude cells are blank. Defined line-endpoint meaning, city/county text and lengths do not supply a qualified center, CRS or accuracy.'],
        'milestones': [
            {'field': 'Original planned in-service (BL)', 'date': original, 'evidence': [source(number, ('BL',))]},
            {'field': 'Actual or expected construction start (BK)', 'date': start, 'evidence': [source(number, ('BK',))]},
        ],
    }
    if kind == 'transmission-project':
        inventory['recordUnit'] = 'component' if parent else 'upgrade'
    return {
        'id': f'{KEY}:{native_id}', 'sourceId': native_id, 'utilityId': KEY, 'sponsor': None,
        'name': text(values['C']), 'description': text(values['G']) or '', 'status': status,
        'voltageKv': voltages(values), 'date': current, 'buildWindow': None, 'location': None,
        'sources': [source(number), letter_source(), definition_source((62, 66))],
        'notes': notes, 'inventory': inventory,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--original-dir', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--scope', choices=('sample', 'full'), default='sample')
    parser.add_argument('--generated-at', default=datetime.now(timezone.utc).isoformat())
    args = parser.parse_args()
    rows, headers = read_rows(args.original_dir)
    retained = [row for row in rows if not exclusion(row['values'])]
    kind_counts = Counter(kind_and_basis(row['values'])[0] for row in retained)
    lifecycle_counts = Counter(lifecycle(row['values']['BI']) for row in retained)
    if kind_counts != EXPECTED_KIND or lifecycle_counts != EXPECTED_LIFECYCLE:
        raise ValueError('Accepted full-source classification/lifecycle boundary changed.')
    physical_units = Counter('component' if parent_id(row['values']) else 'upgrade' for row in retained if kind_and_basis(row['values'])[0] == 'transmission-project')
    if physical_units != {'component': 977, 'upgrade': 24} or not SAMPLE_IDS.issubset({text(row['values']['AI']) for row in retained}):
        raise ValueError('Accepted source work units or diagnostic IDs changed.')
    ledger, projects, exclusions = [], [], Counter()
    for row in rows:
        number, values = row['row'], row['values']
        native_id = text(values['AI'])
        entry = {'sourceId': native_id, 'row': number, 'lineNumber': values['B'], 'source': source(number),
                 'sourceName': text(values['C']), 'rawStatus': values['BI'], 'rawEquipmentClass': text(values['H']),
                 'rawParentId': text(values['AJ']), 'rawCurrentService': text(values['BM'])}
        reason = exclusion(values)
        if reason:
            exclusions[reason] += 1
            entry.update(disposition='excluded', reason=reason)
        else:
            kind, basis = kind_and_basis(values)
            entry.update(recordKind=kind, classificationBasis=basis, lifecycle=lifecycle(values['BI']))
            if args.scope == 'full' or native_id in SAMPLE_IDS:
                project = make_project(row, headers)
                projects.append(project)
                entry.update(disposition='imported', projectId=project['id'])
            else:
                entry.update(disposition='pending', reason='Accepted retained source row; full normalization awaits sample connection.')
        ledger.append(entry)
    imported = 1591 if args.scope == 'full' else 6
    counts = {'sourceRows': len(rows), 'importedRecords': len(projects), 'excludedRows': sum(exclusions.values()),
              'failedRows': 0, 'pendingRows': sum(entry['disposition'] == 'pending' for entry in ledger), 'locatedRecords': 0}
    if counts != {'sourceRows': 2671, 'importedRecords': imported, 'excludedRows': 1080, 'failedRows': 0, 'pendingRows': 1591 - imported, 'locatedRecords': 0}:
        raise ValueError('Accepted complete-source reconciliation changed.')
    limitations = [
        'Public May 1, 2026 Attachment 11 only; no confidential/NDA copy, hidden metadata or person fields are imported.',
        'March 2026 lifecycle evidence and March 10 automated extraction precede publication; no row modification date is established.',
        '2671 source rows reconcile to 1591 retained and 1080 terminal exclusions. The separate 1434 in-progress headline is not reconciled to this row-filter definition.',
        'Full retained scope is 1001 physical transmission components/upgrades, 278 capital context and 312 planning context; programs/allocations have no invented project unit.',
        'Exact native ID #1 tokens preserve case. ID #2 values group work and may carry repeated descriptions/costs; no synthetic parents or name-based deduplication.',
        'Source percentage status bands are milestone-timing estimates, not measured physical completion. Approval fields and past dates do not replace BI status.',
        'BM current projected-or-actual service supplies 1306 day dates and 285 OAI unknowns. Mixed actual/expected starts and original forecasts remain supplementary; construction windows are unknown.',
        'All public coordinate cells are blank. Endpoints are defined but not supplied; city/county text, voltage and length never supply geometry. Asset owners remain unknown.',
        'Costs/ranges retain $000 units and parent/forecast/actual meanings; they are not additive project totals or hypothetical impact inputs.',
        'This utility-limited FERC capital portfolio is not complete California coverage or a unique engineering-project census.',
        'Public factual evidence supports authorized private analysis; no original-workbook publication or open/commercial reuse license is asserted.',
    ]
    if args.scope == 'sample':
        limitations.append('Six diagnostic records are emitted first; 1585 accepted records await provider connection and full normalization.')
    manifest = {
        'sourceKey': KEY, 'editionId': EDITION, 'label': LABEL, 'reportingGroupId': KEY, 'reportingGroupName': GROUP,
        'publisher': 'Pacific Gas and Electric Company', 'processAuthority': 'CPUC Transmission Project Review Process',
        'publishedOn': EDITION, 'retrievedOn': '2026-09-30', 'source': source(), 'sha256': SHA,
        'geographicScope': 'PG&E California public FERC-transmission capital portfolio; source city/county labels do not establish statewide completeness or coordinates.',
        'inventoryScope': 'Entire public CPUC TPR PUBLIC table, rows 5:2675. Retain nonterminal physical components/upgrades and capital/planning context; exclude 1079 Operational and one Cancelled.',
        'scope': args.scope, 'counts': counts,
        'exclusions': [{'reason': reason, 'count': count} for reason, count in sorted(exclusions.items())], 'limitations': limitations,
    }
    dataset = {'schemaVersion': 2, 'mode': 'real', 'title': LABEL + (' — six-record sample' if args.scope == 'sample' else ' — full accepted inventory'),
               'generatedAt': args.generated_at, 'projects': projects, 'sourceEditions': [manifest],
               'coverage': {'processed': len(projects), 'skipped': 1080, 'failed': 0,
                            'notes': [f"All 2671 source rows are reconciled; {counts['pendingRows']} retained entries remain pending.",
                                      'Native source work units are not distinct engineering-project counts; all retained entries remain unlocated.']}}
    quality = {
        'sourceKey': KEY, 'editionId': EDITION, 'scope': args.scope,
        'originals': [{'file': FILE, 'url': URL, 'sha256': SHA}, {'file': LETTER_FILE, 'url': LETTER_URL, 'sha256': LETTER_SHA}],
        'fullSourceStatusCounts': dict(sorted(EXPECTED_STATUS.items())),
        'acceptedRetainedKindCounts': dict(sorted(kind_counts.items())),
        'acceptedRetainedLifecycleCounts': dict(sorted(lifecycle_counts.items())),
        'physicalRecordUnitCounts': dict(sorted(physical_units.items())),
        'contextRecordUnit': 'Omitted for unsupported program/budget units.',
        'acceptedCurrentServicePrecisionCounts': dict(sorted(Counter(milestone(row['values']['BM'], 'Current service')['precision'] for row in retained).items())),
        'importedKindCounts': dict(sorted(Counter(project['inventory']['recordKind'] for project in projects).items())),
        'importedLifecycleCounts': dict(sorted(Counter(project['inventory']['lifecycle'] for project in projects).items())),
        'sampleIds': [f'{KEY}:{native_id}' for native_id in sorted(SAMPLE_IDS)],
        'classificationPrecedence': ['Terminal Operational/Cancelled exclusion', 'Forecast/investment/redaction planning context', 'Bucket/program/right-of-way/supporting capital context', 'Explicit physical equipment transmission component/upgrade'],
        'publicReplacementNamePatterns': list(REPLACEMENT_NAMES),
        'preservedNonTdotParent': {'sourceId': '5744742', 'parentId': 'P.05271', 'source': source(2471, ('AI', 'AJ'))},
        'sourceDates': {'publication': EDITION, 'automatedExtraction': '2026-03-10', 'lifecycleAsOf': '2026-03', 'rowUpdatedOn': None},
        'omittedFields': ['Y Project Manager', 'AD:AH utility prioritization scores', 'CR source cost-benefit score'],
        'locationEvidence': {'allSourceLatitudeBlank': 2671, 'allSourceLongitudeBlank': 2671, 'source': definition_source((7, 8))},
        'headlineDiscrepancy': {'transmittalInProgress': 1434, 'retainedNonterminalRows': 1591, 'meaning': 'Different/unreconciled definitions; no source rows removed to force equality.'},
    }
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for name, value in [('dataset.json', dataset), ('manifest.json', manifest), ('row-ledger.json', ledger), ('source-quality.json', quality)]:
        (args.output_dir / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')
    sizes = {'scope': args.scope, 'datasetRecords': len(projects),
             'compactUtf8Bytes': len(json.dumps(dataset, ensure_ascii=False, separators=(',', ':')).encode()),
             'formattedUtf8Bytes': (args.output_dir / 'dataset.json').stat().st_size}
    (args.output_dir / 'payload-sizes.json').write_text(json.dumps(sizes, indent=2) + '\n')
    print(json.dumps({'counts': counts, 'payloadSizes': sizes}))


if __name__ == '__main__':
    main()
