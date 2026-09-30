#!/usr/bin/env python3
"""Read the ten supplied examples and join their public planning-table records."""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from pypdf import PdfReader

DOMINION_FILE = "2024-2028-2million-and-above-project-descriptions.pdf"
GEORGIA_FILE = "2025 IRP Volume 3 PUBLIC DISCLOSURE.pdf"
# Workbook alias -> source page and native identifier. These are explicit joins,
# never fuzzy matches that could conflate similarly named projects.
DOMINION_JOINS = {
    "DESC_1": (14, "6809 E"),
    "DESC_2": (31, "6810 A"),
    "DESC_3": (23, "06367 D - G"),
    "DESC_4": (1, "6807 B"),
    "DESC_5": (10, "6808 S"),
}
GEORGIA_JOINS = {
    "GPC_1": (189, "20793"),
    "GPC_2": (180, "20277"),
    "GPC_3": (183, "20065"),
    "GPC_4": (177, "18492"),
    "GPC_5": (177, "11821"),
}
TABLE_ROW = re.compile(
    r"(?m)^\s*(\d{3})\s+(20\d{2})\s+(\d{4,6})\s+(.*?)\s+"
    r"(\d{1,2}/\d{1,2}/20\d{2})\s+(\w+)\s+REDACTED", re.S
)


def compact(text: str) -> str:
    return " ".join(text.split())


def source(file: str, locator: str) -> dict:
    return {"file": file, "locator": locator, "url": None}


def stable_id(utility: str, native_id: str) -> str:
    return utility + "-" + re.sub(r"[^a-z0-9]+", "-", native_id.lower()).strip("-")


def date_value(raw: str) -> str:
    for pattern in ("%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(raw.strip(), pattern).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Expected a single day-precision date, received {raw!r}")


def between(text: str, start: str, end: str) -> str:
    return compact(text.split(start, 1)[1].split(end, 1)[0])


def dominion_record(text: str, expected_id: str | None = None) -> dict:
    native_id = between(text, "Project ID", "Project Description")
    if expected_id is not None and native_id != expected_id:
        raise ValueError(f"Dominion join changed: expected {expected_id}, received {native_id}")
    raw_date = between(text, "Planned In-Service Date", "Estimated Project Cost")
    notes = []
    # The public Project Need section is context, not a derived eligibility rule.
    # Omit context if the expected section boundaries are absent or ambiguous.
    needs = re.findall(r"(?ms)^[ \t]*Project Need[ \t]*\n(.*?)^[ \t]*Project Status[ \t]*$", text)
    if len(needs) == 1 and compact(needs[0]):
        notes.append(f"Project need: {compact(needs[0])}")
    try:
        date = {"value": date_value(raw_date), "precision": "day", "kind": "in-service", "label": "Planned in-service date"}
    except ValueError:
        date = {"value": None, "precision": "unknown", "kind": "in-service", "label": "Planned in-service dates by phase" if "phase" in raw_date.lower() else "Planned in-service date"}
        notes.append(f"Source planned in-service dates: {raw_date}. No single date selected.")
    return {
        "sourceId": native_id,
        "sponsor": "Dominion Energy South Carolina",
        "name": between(text, "5 Year Budget", "Project ID"),
        "description": between(text, "Project Description", "Project Need"),
        "status": between(text, "Project Status", "Planned In-Service Date"),
        "date": date,
        "notes": notes,
    }


def georgia_record(text: str, expected_id: str) -> dict:
    rows = [row for row in TABLE_ROW.findall(text) if row[2] == expected_id]
    if len(rows) != 1:
        raise ValueError(f"Expected one Georgia Table 2 row for TEAMS {expected_id}, found {len(rows)}")
    zone, _, native_id, name, need_date, sponsor = rows[0]
    return {
        "sourceId": native_id, "sponsor": sponsor, "name": compact(name),
        "description": "", "status": None,
        "date": {"value": date_value(need_date), "precision": "day", "kind": "need", "label": "Need date"},
        "notes": [f"Planning zone: {zone}."],
    }


def workbook_location(row: dict, number: int) -> dict | None:
    # Read the workbook's cached formula result. Do not reinterpret it as an
    # actual route or independently verified station point.
    lat, lon = row["lat_center"], row["lon_center"]
    if lat is None or lon is None:
        return None
    a_known = row["lat_a"] is not None and row["lon_a"] is not None
    b_known = row["lat_b"] is not None and row["lon_b"] is not None
    if not a_known and not b_known:
        raise ValueError(f"Center has no complete endpoint evidence in projects row {number}")
    return {
        "lat": lat, "lon": lon,
        "method": "endpoint-midpoint" if a_known and b_known else "single-endpoint",
        "confidence": "unverified",
        "evidence": [source("Projects_Overlaps.xlsx", f"projects!E{number}:L{number}")],
    }


def ingest_seed(inputs: Path, generated_at: str) -> dict:
    workbook = load_workbook(inputs / "Projects_Overlaps.xlsx", read_only=True, data_only=True)
    sheet = workbook["projects"]
    values = sheet.iter_rows(values_only=True)
    headers = next(values)
    dominion = PdfReader(inputs / "Project Listings" / "Dominion Energy" / DOMINION_FILE)
    georgia = PdfReader(inputs / "Project Listings" / "Georgia Power" / GEORGIA_FILE)
    projects = []
    for number, values_row in enumerate(values, start=2):
        row = dict(zip(headers, values_row))
        alias = row["project_id"]
        if not alias:
            continue
        if alias in DOMINION_JOINS:
            page, native_id = DOMINION_JOINS[alias]
            record = dominion_record(dominion.pages[page - 1].extract_text(), native_id)
            utility = "desc"
            pdf_source = source(DOMINION_FILE, f"physical page {page}, Project ID {native_id}")
        elif alias in GEORGIA_JOINS:
            page, native_id = GEORGIA_JOINS[alias]
            record = georgia_record(georgia.pages[page - 1].extract_text(), native_id)
            utility = "gpc"
            pdf_source = source(GEORGIA_FILE, f"physical page {page}, Table 2, TEAMS {native_id}")
        else:
            raise ValueError(f"Unmapped workbook alias {alias!r}; add an explicit source join")
        location = workbook_location(row, number)
        notes = record.pop("notes", []) + [f"Workbook alias: {alias}.", "Location supplied in workbook; not independently verified."]
        if location and location["method"] == "single-endpoint":
            notes.append("Only one endpoint is located; the center represents that endpoint, not the whole project.")
        if alias in ("GPC_2", "GPC_3"):
            notes.append("The workbook gives different McIntosh longitudes across these two projects; each source value is retained.")
        projects.append({
            "id": stable_id(utility, record["sourceId"]), "utilityId": utility,
            **record,
            "voltageKv": list(dict.fromkeys(float(value) for value in re.findall(r"(\d+(?:\.\d+)?)\s*kV", record["name"], re.I))),
            "buildWindow": None, "location": location,
            "sources": [pdf_source, source("Projects_Overlaps.xlsx", f"projects!A{number}:M{number}")],
            "notes": notes,
        })
    workbook.close()
    return {
        "schemaVersion": 1, "mode": "real", "title": "Ten supplied planning projects",
        "generatedAt": generated_at, "projects": projects,
        "coverage": {
            "processed": len(projects), "skipped": 0, "failed": 0,
            "notes": [
                "Ten workbook projects linked to public planning records; full filing inventory is not yet included.",
                "Georgia Table 2 need dates retain their source meaning. Dates are historical plan snapshots, not current completion claims.",
                "All supplied locations are unverified centers. Workbook overlap outputs were not imported.",
            ],
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--inputs", type=Path, required=True, help="Preserved Sperry-Tech-Challenge input directory")
    parser.add_argument("--output", type=Path, default=Path("app/src/data/real/dataset.json"))
    parser.add_argument("--generated-at", default=datetime.now(timezone.utc).isoformat(timespec="seconds"))
    args = parser.parse_args()
    dataset = ingest_seed(args.inputs, args.generated_at)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(dataset, indent=2, ensure_ascii=False) + "\n")
    print(f"Wrote {len(dataset['projects'])} real project records to {args.output}")


if __name__ == "__main__":
    main()
